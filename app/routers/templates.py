"""Template Excel: upload, deteksi worksheet, registrasi (pilih log & detail)."""
from __future__ import annotations

import sqlite3
import uuid
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from pydantic import BaseModel

from .. import config, db
from ..deps import audit, current_user, get_db, require_admin

router = APIRouter(prefix="/api/templates", tags=["templates"])


def _inspect_workbook(path: Path, preview_rows: int = 60, preview_cols: int = 26) -> list[dict]:
    """Kembalikan tiap worksheet: sel non-kosong (koordinat+nilai) + merged ranges.

    Dipakai UI Template Mapping untuk memilih anchor sel per kategori foto,
    kolom LOG, dan anchor tabel inventory.
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        cells = []
        maxr = min(ws.max_row or 0, preview_rows)
        maxc = min(ws.max_column or 0, preview_cols)
        for r in range(1, maxr + 1):
            for c in range(1, maxc + 1):
                v = ws.cell(r, c).value
                if v is not None:
                    cells.append({"ref": ws.cell(r, c).coordinate, "text": str(v)[:120]})
        merged = [str(rng) for rng in ws.merged_cells.ranges]
        sheets.append({"name": ws.title, "max_row": ws.max_row or 0,
                       "max_col": ws.max_column or 0, "cells": cells, "merged": merged})
    wb.close()
    return sheets


@router.post("/inspect")
async def inspect_template(file: UploadFile = File(...),
                           conn: sqlite3.Connection = Depends(get_db),
                           user=Depends(require_admin)):
    """Simpan sementara & kembalikan daftar worksheet + preview untuk dipilih."""
    if not (file.filename or "").lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Harus file .xlsx/.xlsm")
    tmp_name = f"pending_{uuid.uuid4().hex}.xlsx"
    tmp_path = config.TEMPLATES_DIR / tmp_name
    tmp_path.write_bytes(await file.read())
    try:
        sheets = _inspect_workbook(tmp_path)
    except Exception as e:
        tmp_path.unlink(missing_ok=True)
        raise HTTPException(400, f"Gagal membaca Excel: {e}")
    return {"pending_file": tmp_name, "orig_name": file.filename, "sheets": sheets}


class RegisterIn(BaseModel):
    pending_file: str
    name: str
    orig_name: str = "template.xlsx"
    sheet_log: str
    sheet_detail: str
    config: dict = {}
    activate: bool = True


@router.post("/register")
def register_template(body: RegisterIn, conn: sqlite3.Connection = Depends(get_db),
                      user=Depends(require_admin)):
    import json
    src = config.TEMPLATES_DIR / body.pending_file
    if not src.exists():
        raise HTTPException(400, "File pending tidak ditemukan, upload ulang")
    final_name = f"tpl_{uuid.uuid4().hex}.xlsx"
    final_path = config.TEMPLATES_DIR / final_name
    src.rename(final_path)
    now = db.now_iso()
    cur = conn.execute(
        "INSERT INTO templates(name,filename,path,sheet_log,sheet_detail,config_json,"
        "active,uploaded_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (body.name, body.orig_name, str(final_path), body.sheet_log, body.sheet_detail,
         json.dumps(body.config, ensure_ascii=False), 1 if body.activate else 0,
         user["id"], now),
    )
    if body.activate:
        conn.execute("UPDATE templates SET active=0 WHERE id<>?", (cur.lastrowid,))
    conn.commit()
    audit(conn, user, "register", "template", cur.lastrowid, body.name)
    return {"ok": True, "id": cur.lastrowid}


@router.get("")
def list_templates(conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    rows = conn.execute(
        "SELECT id,name,filename,sheet_log,sheet_detail,active,created_at "
        "FROM templates ORDER BY id DESC"
    ).fetchall()
    return [dict(r) for r in rows]


@router.post("/{tpl_id}/activate")
def activate_template(tpl_id: int, conn: sqlite3.Connection = Depends(get_db),
                      user=Depends(require_admin)):
    if not conn.execute("SELECT 1 FROM templates WHERE id=?", (tpl_id,)).fetchone():
        raise HTTPException(404, "Template tidak ditemukan")
    conn.execute("UPDATE templates SET active=CASE WHEN id=? THEN 1 ELSE 0 END", (tpl_id,))
    conn.commit()
    audit(conn, user, "activate", "template", tpl_id)
    return {"ok": True}


class RenameIn(BaseModel):
    name: str


@router.patch("/{tpl_id}")
def rename_template(tpl_id: int, body: RenameIn, conn: sqlite3.Connection = Depends(get_db),
                    user=Depends(require_admin)):
    """Ganti nama template terdaftar."""
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "Nama tidak boleh kosong")
    if not conn.execute("SELECT 1 FROM templates WHERE id=?", (tpl_id,)).fetchone():
        raise HTTPException(404, "Template tidak ditemukan")
    conn.execute("UPDATE templates SET name=? WHERE id=?", (name, tpl_id))
    conn.commit()
    audit(conn, user, "rename", "template", tpl_id, name)
    return {"ok": True, "name": name}


@router.delete("/{tpl_id}")
def delete_template(tpl_id: int, conn: sqlite3.Connection = Depends(get_db),
                    user=Depends(require_admin)):
    """Hapus template beserta file-nya. Template aktif tidak boleh dihapus."""
    row = conn.execute("SELECT * FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template tidak ditemukan")
    if row["active"]:
        raise HTTPException(400, "Template aktif tidak bisa dihapus — aktifkan template lain dulu")
    try:
        Path(row["path"]).unlink(missing_ok=True)
    except Exception:
        pass
    conn.execute("DELETE FROM templates WHERE id=?", (tpl_id,))
    conn.commit()
    audit(conn, user, "delete", "template", tpl_id, row["name"])
    return {"ok": True}


@router.get("/{tpl_id}")
def get_template(tpl_id: int, conn: sqlite3.Connection = Depends(get_db),
                 user=Depends(current_user)):
    import json
    row = conn.execute("SELECT * FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template tidak ditemukan")
    d = dict(row)
    d["config"] = json.loads(row["config_json"]) if row["config_json"] else {}
    d.pop("config_json", None)
    return d


@router.get("/{tpl_id}/sheets")
def template_sheets(tpl_id: int, conn: sqlite3.Connection = Depends(get_db),
                    user=Depends(require_admin)):
    """Baca ulang worksheet template terdaftar (untuk UI mapping)."""
    row = conn.execute("SELECT * FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template tidak ditemukan")
    if not Path(row["path"]).exists():
        raise HTTPException(400, "File template hilang di server")
    return {"sheets": _inspect_workbook(Path(row["path"]))}


class MappingIn(BaseModel):
    config: dict            # {log:{...}, detail:{...}} anchor mapping
    sheet_log: str | None = None
    sheet_detail: str | None = None


@router.put("/{tpl_id}/mapping")
def save_mapping(tpl_id: int, body: MappingIn, conn: sqlite3.Connection = Depends(get_db),
                 user=Depends(require_admin)):
    import json
    row = conn.execute("SELECT * FROM templates WHERE id=?", (tpl_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Template tidak ditemukan")
    fields = ["config_json=?"]
    params: list = [json.dumps(body.config, ensure_ascii=False)]
    if body.sheet_log:
        fields.append("sheet_log=?"); params.append(body.sheet_log)
    if body.sheet_detail:
        fields.append("sheet_detail=?"); params.append(body.sheet_detail)
    params.append(tpl_id)
    conn.execute(f"UPDATE templates SET {','.join(fields)} WHERE id=?", params)
    conn.commit()
    audit(conn, user, "save_mapping", "template", tpl_id)
    return {"ok": True}
