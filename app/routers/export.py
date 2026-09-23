"""Generate output: Excel (isi salinan template) & PDF."""
from __future__ import annotations

import json
import re
import sqlite3
import zipfile
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import FileResponse

from .. import config, db
from ..deps import audit, current_user, get_db
from ..services import excel as excel_svc
from ..services import pdf as pdf_svc
from ..services import xlsx2pdf

router = APIRouter(prefix="/api/export", tags=["export"])


def _filter_where(q: str, status: str, creator: int, date: str):
    """Bangun klausa WHERE yang sama dengan GET /api/locations (Log Lokasi)."""
    where, params = [], []
    if q.strip():
        like = f"%{q.strip()}%"
        where.append("(code LIKE ? OR name LIKE ?)")
        params += [like, like]
    if status in ("draft", "selesai"):
        where.append("status = ?")
        params.append(status)
    if creator:
        where.append("created_by = ?")
        params.append(creator)
    if date.strip():
        where.append("date(created_at,'localtime') = ?")
        params.append(date.strip())
    wsql = ("WHERE " + " AND ".join(where)) if where else ""
    return wsql, params


def _gather_locations(conn: sqlite3.Connection, scope: str, loc_id: int | None,
                      q: str = "", status: str = "all", creator: int = 0,
                      date: str = "") -> list[dict]:
    if scope == "all":
        rows = conn.execute("SELECT * FROM locations ORDER BY id").fetchall()
    elif scope == "filter":
        wsql, params = _filter_where(q, status, creator, date)
        rows = conn.execute(f"SELECT * FROM locations {wsql} ORDER BY id", params).fetchall()
    else:
        if not loc_id:
            raise HTTPException(400, "loc_id wajib untuk scope 'one'")
        rows = conn.execute("SELECT * FROM locations WHERE id=?", (loc_id,)).fetchall()
    if not rows:
        raise HTTPException(404, "Tidak ada lokasi untuk diekspor")
    out = []
    for r in rows:
        inv = conn.execute(
            "SELECT nama_barang,merk_type,jumlah,sn_tagging,keterangan FROM inventory_items "
            "WHERE location_id=? ORDER BY sort_order,id", (r["id"],)).fetchall()
        photos = conn.execute(
            "SELECT category,path,ocr_serial FROM photos WHERE location_id=? ORDER BY id",
            (r["id"],)).fetchall()
        out.append({
            "id": r["id"], "code": r["code"], "name": r["name"],
            "data": json.loads(r["data_json"]),
            "inventory": [dict(i) for i in inv],
            "photos": [dict(p) for p in photos],
        })
    return out


def _stamp(prefix: str, ext: str) -> Path:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return config.OUTPUT_DIR / f"{prefix}_{ts}.{ext}"


def _safe_name(s: str) -> str:
    """Nama file aman (untuk entri ZIP)."""
    s = re.sub(r'[\\/:*?"<>|]+', " ", (s or "")).strip()
    return re.sub(r"\s+", " ", s) or "lokasi"


def _active_template(conn: sqlite3.Connection):
    """Template aktif + config, atau (None, None) bila tak ada / file hilang."""
    tpl = conn.execute("SELECT * FROM templates WHERE active=1 ORDER BY id DESC LIMIT 1").fetchone()
    if not tpl or not Path(tpl["path"]).exists():
        return None, None
    tcfg = json.loads(tpl["config_json"]) if tpl["config_json"] else {}
    tcfg["sheet_log"] = tpl["sheet_log"]
    tcfg["sheet_detail"] = tpl["sheet_detail"]
    return tpl, tcfg


def _template_pdf(tpl, tcfg, locs, out_pdf: Path) -> bool:
    """Isi template -> xlsx (detail saja, tanpa LOG) -> konversi PDF (Excel/LO).
    True bila berhasil; False agar pemanggil fallback ke reportlab."""
    if not xlsx2pdf.available():
        return False
    import openpyxl
    tmp_xlsx = out_pdf.with_suffix(".xlsx")
    ok = False
    try:
        excel_svc.build_workbook(tpl["path"], tcfg, locs, str(tmp_xlsx))
        wb = openpyxl.load_workbook(str(tmp_xlsx))
        log_name = tcfg.get("sheet_log")
        if log_name and log_name in wb.sheetnames and len(wb.sheetnames) > 1:
            del wb[log_name]                  # PDF hanya halaman detail lokasi
        wb.save(str(tmp_xlsx))
        ok = xlsx2pdf.xlsx_to_pdf(str(tmp_xlsx), str(out_pdf))
    except Exception:
        ok = False
    finally:
        try:
            tmp_xlsx.unlink()
        except OSError:
            pass
    return bool(ok) and out_pdf.exists()


def _render_pdf(conn, locs, out_pdf: Path, cats, title) -> str:
    """PDF mengikuti template aktif bila memungkinkan; jika tidak, pakai
    generator bawaan (reportlab). Mengembalikan 'template' atau 'builtin'."""
    tpl, tcfg = _active_template(conn)
    if tpl and _template_pdf(tpl, tcfg, locs, out_pdf):
        return "template"
    pdf_svc.build_pdf(locs, cats, str(out_pdf), app_title=title)
    return "builtin"


@router.get("/excel")
def export_excel(scope: str = Query("one"), loc_id: int | None = None,
                 q: str = Query(""), status: str = Query("all"),
                 creator: int = Query(0), date: str = Query(""),
                 conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    tpl = conn.execute("SELECT * FROM templates WHERE active=1 ORDER BY id DESC LIMIT 1").fetchone()
    if not tpl:
        raise HTTPException(400, "Belum ada template aktif. Daftarkan template dulu di menu Template.")
    if not Path(tpl["path"]).exists():
        raise HTTPException(400, "File template hilang di server.")
    locs = _gather_locations(conn, scope, loc_id, q, status, creator, date)
    tcfg = json.loads(tpl["config_json"]) if tpl["config_json"] else {}
    tcfg["sheet_log"] = tpl["sheet_log"]
    tcfg["sheet_detail"] = tpl["sheet_detail"]
    prefix = locs[0]["code"] if scope == "one" else "Log_BAA"
    out = _stamp(prefix, "xlsx")
    res = excel_svc.build_workbook(tpl["path"], tcfg, locs, str(out))
    audit(conn, user, "export_excel", "export", scope, out.name)
    resp = FileResponse(str(out), filename=out.name,
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    if res.get("warnings"):
        resp.headers["X-Export-Warnings"] = str(len(res["warnings"]))
    return resp


@router.get("/pdf")
def export_pdf(scope: str = Query("one"), loc_id: int | None = None,
               q: str = Query(""), status: str = Query("all"),
               creator: int = Query(0), date: str = Query(""),
               conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    locs = _gather_locations(conn, scope, loc_id, q, status, creator, date)
    cats = db.get_setting(conn, "photo_categories", [])
    title = db.get_setting(conn, "app_title", "Berita Acara Aktivasi")
    out = _stamp(locs[0]["code"] if scope == "one" else "BAA", "pdf")
    mode = _render_pdf(conn, locs, out, cats, title)
    audit(conn, user, "export_pdf", "export", scope, out.name)
    resp = FileResponse(str(out), filename=out.name, media_type="application/pdf")
    resp.headers["X-PDF-Mode"] = mode
    return resp


@router.get("/pdf-zip")
def export_pdf_zip(scope: str = Query("filter"), loc_id: int | None = None,
                   q: str = Query(""), status: str = Query("all"),
                   creator: int = Query(0), date: str = Query(""),
                   conn: sqlite3.Connection = Depends(get_db), user=Depends(current_user)):
    """Satu PDF detail per lokasi (tanpa LOG sheet), dibundel dalam satu ZIP."""
    locs = _gather_locations(conn, scope, loc_id, q, status, creator, date)
    cats = db.get_setting(conn, "photo_categories", [])
    title = db.get_setting(conn, "app_title", "Berita Acara Aktivasi")
    zpath = _stamp("PDF_BAA", "zip")
    used: set[str] = set()
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED) as zf:
        for loc in locs:
            nm = _safe_name(f"{loc.get('code','')} {loc.get('name','') or loc.get('data',{}).get('nama_lokasi','')}")
            entry = f"{nm}.pdf"
            i = 2
            while entry.lower() in used:
                entry = f"{nm} ({i}).pdf"; i += 1
            used.add(entry.lower())
            tmp = config.OUTPUT_DIR / f"_tmp_{loc['id']}.pdf"
            _render_pdf(conn, [loc], tmp, cats, title)
            zf.write(str(tmp), entry)
            try:
                tmp.unlink()
            except OSError:
                pass
    audit(conn, user, "export_pdf_zip", "export", scope, zpath.name)
    return FileResponse(str(zpath), filename=zpath.name, media_type="application/zip")
