"""Tulis data ke SALINAN template Excel (format asli dipertahankan).

Pendekatan config-driven: mapping kolom LOG, anchor tabel inventory, dan anchor
foto per-kategori diambil dari ``template.config_json`` (fallback ke DEFAULT yang
diturunkan dari Template_BAA.xlsx). Anchor foto bisa dikalibrasi lewat Pengaturan
tanpa mengubah kode.
"""
from __future__ import annotations

import copy
from pathlib import Path

# Default mapping untuk Template_BAA.xlsx (bisa dioverride per template)
DEFAULT_CONFIG = {
    "log": {
        "start_row": 3,
        "columns": {
            "no": "A", "lokasi": "B", "nama_barang": "C", "merk_type": "D",
            "jumlah": "E", "sn_tagging": "F", "keterangan": "G", "foto_lengkap": "H",
        },
    },
    "detail": {
        # sel judul lokasi (opsional) — diisi nama lokasi
        "location_cells": {},
        "inventory": {
            "start_row": 32, "max_rows": 5,
            "cols": {"nama_barang": "B", "merk_type": "C", "jumlah": "E",
                     "sn_tagging": "F", "keterangan": "H"},
        },
        # anchor sel kiri-atas tiap kategori foto (default utk Template_BAA;
        # sesuaikan di menu Mapping bila template berbeda)
        "photos": {
            "dashboard": "B27", "tampak_depan": "B74", "teknisi": "B91",
            "outdoor": "K26", "indoor": "K43", "sn_kit": "K74", "sn_router": "K91",
            "sn_ap": "T24", "ping": "T43", "speed": "T74", "simkopdes": "T91",
        },
        "photo_max_w": 320, "photo_max_h": 240,
    },
}


def _merge(base: dict, override: dict) -> dict:
    out = copy.deepcopy(base)
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = _merge(out[k], v)
        else:
            out[k] = v
    return out


def _fit(img_w: int, img_h: int, max_w: int, max_h: int) -> tuple[int, int]:
    if img_w <= 0 or img_h <= 0:
        return max_w, max_h
    r = min(max_w / img_w, max_h / img_h, 1.0)
    return int(img_w * r), int(img_h * r)


# Karakter yang dilarang Excel untuk nama worksheet
_BAD_SHEET = set(r':\/?*[]')


def _safe_sheet_title(code: str, name: str, used: set) -> str:
    """Nama worksheet = 'Kode Nama Lokasi', dibersihkan & dipotong <=31 char,
    dan dijamin unik dalam workbook."""
    raw = f"{code} {name}".strip() if name else (code or "Lokasi")
    clean = "".join(" " if ch in _BAD_SHEET else ch for ch in raw)
    clean = " ".join(clean.split()).strip("'").strip() or (code or "Lokasi")
    base = clean[:31]
    title = base
    n = 2
    while title in used or title == "":
        suffix = f" ({n})"
        title = base[:31 - len(suffix)].rstrip() + suffix
        n += 1
    used.add(title)
    return title


def _letterbox(path: str, box_w: int, box_h: int):
    """Kembalikan BytesIO PNG berukuran persis box_w x box_h: foto di-fit
    (tanpa distorsi & tanpa dipotong) lalu ditaruh di tengah kanvas putih.
    Membuat semua foto hasil export punya ukuran yang seragam."""
    import io
    from PIL import Image as PILImage
    im = PILImage.open(path)
    if im.mode not in ("RGB",):
        im = im.convert("RGB")
    im.thumbnail((box_w, box_h), PILImage.LANCZOS)   # hanya mengecilkan, jaga rasio
    canvas = PILImage.new("RGB", (box_w, box_h), (255, 255, 255))
    canvas.paste(im, ((box_w - im.width) // 2, (box_h - im.height) // 2))
    bio = io.BytesIO()
    canvas.save(bio, format="PNG")
    bio.seek(0)
    return bio


def _copy_print_settings(src, dst) -> None:
    """openpyxl copy_worksheet() tidak menyalin print area & pengaturan cetak.
    Salin manual agar sheet hasil tetap punya Print Area seperti template."""
    import copy
    try:
        if src.print_area:
            dst.print_area = src.print_area
    except Exception:
        pass
    for attr in ("page_setup", "page_margins", "print_options",
                 "sheet_properties", "print_title_rows", "print_title_cols"):
        try:
            val = getattr(src, attr)
            if val is None:
                continue
            setattr(dst, attr, copy.copy(val) if not isinstance(val, str) else val)
        except Exception:
            pass


# Perkiraan standar Excel bila lebar kolom / tinggi baris tidak diset eksplisit
_DEF_COL_CHARS = 8.43
_DEF_ROW_PT = 15.0


_MDW_CACHE: dict = {}
# Maximum Digit Width (px) perkiraan bila font tak bisa diukur langsung.
_MDW_LOOKUP = {
    "calibri": 7.0, "aptos": 7.0, "aptos narrow": 6.0, "arial": 7.0,
    "arial narrow": 6.0, "times new roman": 7.0, "verdana": 8.0,
    "tahoma": 7.0, "segoe ui": 7.0,
}


def _mdw_px(font_name: str | None, size) -> float:
    """Lebar digit maksimum (px @96dpi) untuk font+ukuran tertentu.
    Diukur langsung via PIL (akurat di mesin yang punya fontnya), fallback tabel."""
    name = (font_name or "Calibri").strip()
    try:
        size = float(size or 11.0)
    except Exception:
        size = 11.0
    key = (name.lower(), round(size, 1))
    if key in _MDW_CACHE:
        return _MDW_CACHE[key]
    val = None
    try:
        from PIL import ImageFont
        px = int(round(size * 96.0 / 72.0))
        for cand in (name, name + " Regular", name.replace(" ", ""), name.replace(" ", "") + "-Regular"):
            for ext in ("", ".ttf", ".ttc", ".otf"):
                try:
                    fnt = ImageFont.truetype(cand + ext, px)
                    val = max(fnt.getlength(str(d)) for d in range(10))
                    break
                except Exception:
                    continue
            if val:
                break
    except Exception:
        val = None
    if not val:
        val = _MDW_LOOKUP.get(key[0], 7.0)
    _MDW_CACHE[key] = val
    return val


def _col_px(ws, col_idx: int, mdw: float = 7.0) -> int:
    """Lebar kolom (1-based) dalam piksel, memperhitungkan MDW font aktif."""
    from openpyxl.utils import get_column_letter
    dim = ws.column_dimensions.get(get_column_letter(col_idx))
    w = dim.width if (dim is not None and dim.width) else None
    if not w:
        w = ws.sheet_format.defaultColWidth
        if not w:
            w = (ws.sheet_format.baseColWidth or 8) + 0.43   # perkiraan default Excel
    return int(round(w * mdw)) + 5        # rumus lebar Excel (MDW sesuai font)


def _row_px(ws, row_idx: int) -> int:
    """Tinggi baris (1-based) dalam piksel."""
    from openpyxl.utils.units import points_to_pixels
    dim = ws.row_dimensions.get(row_idx)
    h = dim.height if (dim is not None and dim.height) else None
    if not h:
        h = (ws.sheet_format.defaultRowHeight or _DEF_ROW_PT)
    return int(round(points_to_pixels(h)))


def _place_photo(ws, xi, anchor: str) -> None:
    """Tempel gambar. Bila sel anchor adalah bagian dari cell yang di-merge,
    gambar diletakkan di TENGAH (horizontal + vertikal) area merge. Bila bukan
    merge, perilaku tetap seperti biasa (pojok kiri-atas sel anchor)."""
    from openpyxl.utils.cell import coordinate_to_tuple
    from openpyxl.utils.units import pixels_to_EMU
    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
    from openpyxl.drawing.xdr import XDRPositiveSize2D

    row, col = coordinate_to_tuple(anchor)
    rng = None
    for m in ws.merged_cells.ranges:
        if m.min_row <= row <= m.max_row and m.min_col <= col <= m.max_col:
            rng = m
            break
    if rng is None:                        # bukan cell merge -> biasa
        ws.add_image(xi, anchor)
        return

    try:
        f = ws.cell(row=row, column=col).font
        mdw = _mdw_px(f.name, f.sz)
    except Exception:
        mdw = 7.0
    area_w = sum(_col_px(ws, c, mdw) for c in range(rng.min_col, rng.max_col + 1))
    area_h = sum(_row_px(ws, r) for r in range(rng.min_row, rng.max_row + 1))
    # Muatkan gambar ke dalam area merge (skala turun bila lebih besar), sisakan
    # sedikit margin, agar bisa benar-benar center vertikal + horizontal.
    img_w, img_h = int(xi.width), int(xi.height)
    avail_w, avail_h = max(1, area_w - 8), max(1, area_h - 8)
    if img_w > avail_w or img_h > avail_h:
        scale = min(avail_w / img_w, avail_h / img_h)
        img_w = max(1, int(img_w * scale))
        img_h = max(1, int(img_h * scale))
        xi.width, xi.height = img_w, img_h
    off_x = max(0, (area_w - img_w) // 2)
    off_y = max(0, (area_h - img_h) // 2)
    marker = AnchorMarker(col=rng.min_col - 1, colOff=pixels_to_EMU(off_x),
                          row=rng.min_row - 1, rowOff=pixels_to_EMU(off_y))
    xi.anchor = OneCellAnchor(
        _from=marker,
        ext=XDRPositiveSize2D(pixels_to_EMU(int(xi.width)), pixels_to_EMU(int(xi.height))),
    )
    ws.add_image(xi)


def build_workbook(template_path: str, template_config: dict, locations: list[dict],
                   out_path: str) -> dict:
    """locations: list of dict {code,name,data,inventory:[...],photos:[{category,path,...}]}"""
    import openpyxl
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.utils import column_index_from_string

    cfg = _merge(DEFAULT_CONFIG, template_config or {})
    wb = openpyxl.load_workbook(template_path)

    # Nama sheet log & detail dari config_json (disimpan saat registrasi)
    sheet_log = template_config.get("sheet_log") if template_config else None
    sheet_detail = template_config.get("sheet_detail") if template_config else None
    sheet_log = sheet_log or ("LOG" if "LOG" in wb.sheetnames else wb.sheetnames[0])
    sheet_detail = sheet_detail or (wb.sheetnames[1] if len(wb.sheetnames) > 1 else wb.sheetnames[0])

    log_ws = wb[sheet_log]
    detail_tpl = wb[sheet_detail]
    # Ganti nama sheet template detail agar salinan tak bentrok namanya
    _tpl_title = "__tpl_detail__"
    detail_tpl.title = _tpl_title

    warnings = []
    log_row = cfg["log"]["start_row"]
    log_cols = cfg["log"]["columns"]
    used_titles: set = set()
    _img_keep: list = []          # tahan buffer gambar sampai workbook disimpan

    for i, loc in enumerate(locations, start=1):
        data = loc.get("data", {})
        inv = loc.get("inventory", [])
        photos = loc.get("photos", [])
        nama = data.get("nama_lokasi") or loc.get("name") or loc.get("code")

        # --- baris LOG: satu baris per item inventory (list semua) ---
        def _set(row, col_key, value):
            col = log_cols.get(col_key)
            if col:
                log_ws[f"{col}{row}"] = value

        rows_inv = inv if inv else [{}]          # lokasi tanpa inventory tetap 1 baris
        for j, item in enumerate(rows_inv):
            if j == 0:                           # kolom identitas hanya di baris pertama
                _set(log_row, "no", i)
                _set(log_row, "lokasi", nama)
                _set(log_row, "foto_lengkap", "Ya" if len(photos) >= 1 else "Belum")
            _set(log_row, "nama_barang", item.get("nama_barang", ""))
            _set(log_row, "merk_type", item.get("merk_type", ""))
            _set(log_row, "jumlah", item.get("jumlah", ""))
            _set(log_row, "sn_tagging", item.get("sn_tagging", ""))
            _set(log_row, "keterangan", item.get("keterangan", ""))
            log_row += 1

        # --- sheet detail per lokasi (duplikasi template) ---
        ws = wb.copy_worksheet(detail_tpl)
        nama_only = data.get("nama_lokasi") or loc.get("name") or ""   # tanpa fallback kode
        ws.title = _safe_sheet_title(loc.get("code", f"Lokasi_{i:04d}"), nama_only, used_titles)
        _copy_print_settings(detail_tpl, ws)   # copy_worksheet tak menyalin print area/page setup

        for cell, field in cfg["detail"].get("location_cells", {}).items():
            try:
                ws[cell] = data.get(field, "")
            except Exception:
                pass

        invc = cfg["detail"]["inventory"]
        r0 = invc["start_row"]
        for r_off, item in enumerate(inv[: invc.get("max_rows", 999)]):
            r = r0 + r_off
            for field, col in invc["cols"].items():
                try:
                    ws[f"{col}{r}"] = item.get(field, "")
                except Exception:
                    pass

        # --- foto per kategori ---
        photo_anchors = cfg["detail"].get("photos", {})
        by_cat: dict[str, list] = {}
        for p in photos:
            by_cat.setdefault(p.get("category", "uncategorized"), []).append(p)
        for cat, plist in by_cat.items():
            anchor = photo_anchors.get(cat)
            if not anchor or not plist:
                if not anchor:
                    warnings.append(f"{ws.title}: anchor foto '{cat}' belum dikalibrasi")
                continue
            p = plist[0]
            path = p.get("path")
            if not path or not Path(path).exists():
                continue
            box_w = cfg["detail"]["photo_max_w"]
            box_h = cfg["detail"]["photo_max_h"]
            try:
                # Letterbox -> semua foto berukuran seragam (box_w x box_h)
                try:
                    src = _letterbox(path, box_w, box_h)
                    xi = XLImage(src)
                    _img_keep.append(src)          # jaga buffer sampai wb.save()
                except Exception:
                    xi = XLImage(path)             # fallback: foto asli
                xi.width, xi.height = box_w, box_h
                _place_photo(ws, xi, anchor)
            except Exception as e:
                warnings.append(f"{ws.title}: gagal sisip foto '{cat}': {e}")

    # Hapus sheet template detail (yang sudah diganti nama) bila sudah ada salinan
    if locations and _tpl_title in wb.sheetnames and len(wb.sheetnames) > 1:
        try:
            del wb[_tpl_title]
        except Exception:
            pass
    else:
        # tak ada lokasi -> kembalikan nama sheet detail seperti semula
        wb[_tpl_title].title = sheet_detail

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return {"path": out_path, "warnings": warnings}
